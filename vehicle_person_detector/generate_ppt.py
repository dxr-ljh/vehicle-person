import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

base_dir = os.path.dirname(os.path.abspath(__file__))
output_path = os.path.join(base_dir, "汽车与行人检测系统_演示PPT.pptx")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BLUE = RGBColor(0x1E, 0x88, 0xE5)
DARK_BLUE = RGBColor(0x0D, 0x47, 0xA1)
LIGHT_BLUE = RGBColor(0xBB, 0xDE, 0xFB)
DARK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT_GRAY = RGBColor(0xF5, 0xF5, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x43, 0xA0, 0x47)
RED = RGBColor(0xE5, 0x39, 0x35)
ORANGE = RGBColor(0xFF, 0x98, 0x00)


def add_bg(slide, color=WHITE):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_text(slide, text, left, top, width, height, font_size=18,
             color=DARK, bold=False, align=PP_ALIGN.LEFT, font_name="微软雅黑"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = align
    return txBox


def add_title_bar(slide, title_text):
    add_rect(slide, Inches(0), Inches(0), prs.slide_width, Inches(0.1), BLUE)
    add_rect(slide, Inches(0.6), Inches(0.5), Inches(0.08), Inches(0.6), BLUE)
    add_text(slide, title_text, Inches(0.85), Inches(0.5), Inches(10), Inches(0.6),
             font_size=32, color=DARK, bold=True)


def add_bullet_list(slide, items, left, top, width, height, font_size=20, color=DARK):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = "●  " + item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = "微软雅黑"
        p.space_after = Pt(10)
    return txBox


# ==================== 第1页：封面 ====================
slide1 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide1, WHITE)

add_rect(slide1, Inches(0), Inches(0), prs.slide_width, Inches(3.5), BLUE)
add_rect(slide1, Inches(0), Inches(3.5), prs.slide_width, Inches(0.08), DARK_BLUE)

add_text(slide1, "汽车与行人检测系统", Inches(1), Inches(1.2), Inches(11), Inches(1.2),
         font_size=54, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

add_text(slide1, "基于 YOLOv8 的目标检测系统", Inches(1), Inches(2.4), Inches(11), Inches(0.8),
         font_size=28, color=LIGHT_BLUE, align=PP_ALIGN.CENTER)

add_text(slide1, "PyTorch  ·  Ultralytics  ·  OpenCV", Inches(1), Inches(3.0), Inches(11), Inches(0.5),
         font_size=18, color=LIGHT_BLUE, align=PP_ALIGN.CENTER)

add_text(slide1, "演示人：XXX", Inches(1), Inches(5.5), Inches(11), Inches(0.6),
         font_size=22, color=GRAY, align=PP_ALIGN.CENTER)

add_text(slide1, "2026年10月", Inches(1), Inches(6.2), Inches(11), Inches(0.5),
         font_size=18, color=GRAY, align=PP_ALIGN.CENTER)


# ==================== 第2页：项目背景 ====================
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide2, WHITE)
add_title_bar(slide2, "项目背景")

add_text(slide2, "为什么需要汽车与行人检测？", Inches(0.85), Inches(1.6), Inches(11), Inches(0.6),
         font_size=24, color=BLUE, bold=True)

items_left = [
    "智能交通：车流量统计、拥堵监测",
    "自动驾驶：障碍物识别、行人避让",
    "安防监控：异常行为检测、人员追踪",
    "智慧城市：人流密度分析、公共安全",
]

items_right = [
    "传统方法：手工特征提取，准确率低",
    "深度学习：端到端学习，鲁棒性强",
    "实时性要求：视频检测需高帧率处理",
    "多目标检测：同时识别多种类别目标",
]

add_bullet_list(slide2, items_left, Inches(1), Inches(2.5), Inches(5.5), Inches(4),
                font_size=20, color=DARK)
add_bullet_list(slide2, items_right, Inches(7), Inches(2.5), Inches(5.5), Inches(4),
                font_size=20, color=DARK)


# ==================== 第3页：技术方案 ====================
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide3, WHITE)
add_title_bar(slide3, "技术方案")

add_text(slide3, "核心技术选型", Inches(0.85), Inches(1.6), Inches(11), Inches(0.6),
         font_size=24, color=BLUE, bold=True)

techs = [
    ("YOLOv8", "Ultralytics 出品\n单阶段目标检测算法\n速度快、精度高", BLUE),
    ("PyTorch", "Meta 开源深度学习框架\n动态计算图\n生态丰富", ORANGE),
    ("OpenCV", "计算机视觉库\n图像/视频处理\n实时性能好", GREEN),
    ("Tkinter", "Python 内置 GUI\n轻量级界面开发\n跨平台支持", RED),
]

card_w = Inches(2.7)
card_h = Inches(3.5)
start_x = Inches(0.85)
gap = Inches(0.35)

for i, (name, desc, color) in enumerate(techs):
    x = start_x + (card_w + gap) * i
    y = Inches(2.5)

    add_rect(slide3, x, y, card_w, card_h, LIGHT_GRAY)
    add_rect(slide3, x, y, card_w, Inches(0.08), color)

    add_text(slide3, name, x, y + Inches(0.5), card_w, Inches(0.8),
             font_size=28, color=color, bold=True, align=PP_ALIGN.CENTER)

    txBox = slide3.shapes.add_textbox(x + Inches(0.2), y + Inches(1.4),
                                       card_w - Inches(0.4), Inches(2))
    tf = txBox.text_frame
    tf.word_wrap = True
    lines = desc.split("\n")
    for j, line in enumerate(lines):
        if j == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(16)
        p.font.color.rgb = GRAY
        p.font.name = "微软雅黑"
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(8)


# ==================== 第4页：系统架构 ====================
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide4, WHITE)
add_title_bar(slide4, "系统架构")

add_text(slide4, "系统整体流程", Inches(0.85), Inches(1.6), Inches(11), Inches(0.6),
         font_size=24, color=BLUE, bold=True)

layers = [
    ("输入层", "图片 / 视频 / 实时摄像头", ORANGE),
    ("预处理层", "图像缩放、归一化、Tensor 转换", GREEN),
    ("检测层", "YOLOv8 模型推理\n特征提取 → 目标定位 → 类别分类", BLUE),
    ("后处理层", "NMS 非极大值抑制\n置信度过滤、坐标转换", RED),
    ("输出层", "绘制检测框 / 统计信息 / 结果保存", DARK),
]

layer_h = Inches(0.85)
start_y = Inches(2.4)
gap = Inches(0.15)

for i, (name, desc, color) in enumerate(layers):
    y = start_y + (layer_h + gap) * i
    x = Inches(2)
    w = Inches(9.3)

    add_rect(slide4, x, y, Inches(1.8), layer_h, color)
    add_text(slide4, name, x, y, Inches(1.8), layer_h,
             font_size=20, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide4, "", x, y + Inches(0.3), Inches(1.8), Inches(0.4),
             font_size=12, color=WHITE, align=PP_ALIGN.CENTER)

    add_rect(slide4, x + Inches(1.8), y, w - Inches(1.8), layer_h, LIGHT_GRAY)
    add_text(slide4, desc, x + Inches(2), y + Inches(0.05), w - Inches(2.2), layer_h,
             font_size=18, color=DARK, align=PP_ALIGN.LEFT)

    if i < len(layers) - 1:
        arrow_y = y + layer_h + Inches(0.01)
        add_text(slide4, "▼", Inches(6.4), arrow_y, Inches(0.5), Inches(0.15),
                 font_size=14, color=GRAY, align=PP_ALIGN.CENTER)


# ==================== 第5页：功能展示 - 图片检测 ====================
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide5, WHITE)
add_title_bar(slide5, "功能展示 — 图片检测")

left_x = Inches(0.85)
right_x = Inches(7.2)
y = Inches(1.7)

img_path = os.path.join(base_dir, "output", "result_demo_yolo_bus.jpg")
if os.path.exists(img_path):
    slide5.shapes.add_picture(img_path, left_x, y, width=Inches(6), height=Inches(4.5))
else:
    add_rect(slide5, left_x, y, Inches(6), Inches(4.5), LIGHT_GRAY)
    add_text(slide5, "[检测效果图]", left_x, y + Inches(2), Inches(6), Inches(0.5),
             font_size=20, color=GRAY, align=PP_ALIGN.CENTER)

add_text(slide5, "图片检测功能", right_x, y, Inches(5.3), Inches(0.6),
         font_size=24, color=BLUE, bold=True)

img_features = [
    "支持 JPG / PNG / BMP 等常见格式",
    "自动缩放至模型输入尺寸 (640×640)",
    "检测框 + 中文标签 + 置信度显示",
    "实时统计行人和汽车数量",
    "置信度阈值可调节 (0.1 ~ 0.9)",
    "检测结果可保存为图片文件",
]

add_bullet_list(slide5, img_features, right_x, y + Inches(0.8), Inches(5.3), Inches(4),
                font_size=18, color=DARK)


# ==================== 第6页：功能展示 - 视频检测 ====================
slide6 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide6, WHITE)
add_title_bar(slide6, "功能展示 — 视频检测")

left_x = Inches(0.85)
right_x = Inches(7.2)
y = Inches(1.7)

add_rect(slide6, left_x, y, Inches(6), Inches(4.5), DARK)
add_text(slide6, "🎬 视频检测演示", left_x, y + Inches(0.3), Inches(6), Inches(0.6),
         font_size=22, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(slide6, "逐帧检测 + 结果输出", left_x, y + Inches(1.0), Inches(6), Inches(0.5),
         font_size=16, color=LIGHT_BLUE, align=PP_ALIGN.CENTER)

stats = [
    ("支持格式", "MP4 / AVI / MOV / MKV"),
    ("处理方式", "逐帧检测，输出带标注视频"),
    ("统计信息", "累计检测、单帧最多、平均值"),
    ("进度显示", "实时进度条 + 帧计数"),
]

stat_y = y + Inches(2.0)
for i, (label, value) in enumerate(stats):
    row_y = stat_y + Inches(0.55) * i
    add_rect(slide6, left_x + Inches(0.5), row_y, Inches(1.5), Inches(0.45), BLUE)
    add_text(slide6, label, left_x + Inches(0.5), row_y + Inches(0.05), Inches(1.5), Inches(0.4),
             font_size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_rect(slide6, left_x + Inches(2.0), row_y, Inches(3.5), Inches(0.45), LIGHT_GRAY)
    add_text(slide6, value, left_x + Inches(2.1), row_y + Inches(0.05), Inches(3.3), Inches(0.4),
             font_size=14, color=DARK, align=PP_ALIGN.LEFT)

add_text(slide6, "视频检测功能", right_x, y, Inches(5.3), Inches(0.6),
         font_size=24, color=BLUE, bold=True)

vid_features = [
    "支持多种视频格式输入",
    "逐帧 YOLOv8 目标检测",
    "绘制检测框并保存结果视频",
    "实时进度条显示处理进度",
    "统计行人/车辆累计数量",
    "输出视频保留原始帧率和分辨率",
]

add_bullet_list(slide6, vid_features, right_x, y + Inches(0.8), Inches(5.3), Inches(4),
                font_size=18, color=DARK)


# ==================== 第7页：检测效果 ====================
slide7 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide7, WHITE)
add_title_bar(slide7, "检测效果展示")

add_text(slide7, "经典 YOLO 测试图 — 公交车场景", Inches(0.85), Inches(1.5), Inches(11), Inches(0.5),
         font_size=22, color=BLUE, bold=True)

img_path = os.path.join(base_dir, "output", "result_demo_yolo_bus.jpg")
if os.path.exists(img_path):
    slide7.shapes.add_picture(img_path, Inches(0.85), Inches(2.2),
                              width=Inches(7.5), height=Inches(5))
else:
    add_rect(slide7, Inches(0.85), Inches(2.2), Inches(7.5), Inches(5), LIGHT_GRAY)
    add_text(slide7, "[检测效果图]", Inches(0.85), Inches(4), Inches(7.5), Inches(0.6),
             font_size=22, color=GRAY, align=PP_ALIGN.CENTER)

result_x = Inches(8.8)
result_y = Inches(2.2)

add_rect(slide7, result_x, result_y, Inches(3.8), Inches(1.5), LIGHT_BLUE)
add_text(slide7, "检测结果统计", result_x, result_y + Inches(0.15), Inches(3.8), Inches(0.5),
         font_size=20, color=DARK_BLUE, bold=True, align=PP_ALIGN.CENTER)

add_text(slide7, "👤 行人：3 个", result_x, result_y + Inches(0.7), Inches(3.8), Inches(0.4),
         font_size=22, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
add_text(slide7, "🚗 汽车：1 辆", result_x, result_y + Inches(1.05), Inches(3.8), Inches(0.4),
         font_size=22, color=RED, bold=True, align=PP_ALIGN.CENTER)

metric_y = result_y + Inches(1.8)
metrics = [
    ("模型", "YOLOv8n (nano)"),
    ("置信度阈值", "0.3"),
    ("推理速度", "~50ms/帧 (CPU)"),
    ("输入尺寸", "640 × 640"),
    ("检测类别", "行人 / 汽车 / 公交 / 卡车"),
]

for i, (label, value) in enumerate(metrics):
    row_y = metric_y + Inches(0.55) * i
    add_rect(slide7, result_x, row_y, Inches(1.6), Inches(0.45), BLUE)
    add_text(slide7, label, result_x, row_y + Inches(0.05), Inches(1.6), Inches(0.4),
             font_size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_rect(slide7, result_x + Inches(1.6), row_y, Inches(2.2), Inches(0.45), LIGHT_GRAY)
    add_text(slide7, value, result_x + Inches(1.7), row_y + Inches(0.05), Inches(2), Inches(0.4),
             font_size=14, color=DARK, align=PP_ALIGN.LEFT)


# ==================== 第8页：系统界面 ====================
slide8 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide8, WHITE)
add_title_bar(slide8, "图形操作界面")

add_text(slide8, "基于 Tkinter 的图形化操作界面", Inches(0.85), Inches(1.5), Inches(11), Inches(0.5),
         font_size=22, color=BLUE, bold=True)

add_rect(slide8, Inches(0.85), Inches(2.2), Inches(11.6), Inches(4.8), LIGHT_GRAY)
add_rect(slide8, Inches(0.85), Inches(2.2), Inches(8.5), Inches(4.8), RGBColor(0x1E, 0x1E, 0x1E))
add_text(slide8, "🖼  图片显示区域", Inches(0.85), Inches(4.2), Inches(8.5), Inches(0.8),
         font_size=24, color=RGBColor(0xAA, 0xAA, 0xAA), bold=True, align=PP_ALIGN.CENTER)

panel_x = Inches(9.5)
panel_y = Inches(2.2)
panel_w = Inches(2.95)

sections = [
    ("🖼 图片检测", ["选择图片", "生成测试图", "🔍 检测图片"]),
    ("🎬 视频检测", ["选择视频", "生成测试视频", "▶️ 检测视频"]),
    ("⚙️ 设置", ["置信度滑块"]),
    ("📊 检测结果", ["数量统计", "详细列表"]),
]

sy = panel_y
for title, btns in sections:
    add_rect(slide8, panel_x, sy, panel_w, Inches(0.4), BLUE)
    add_text(slide8, title, panel_x + Inches(0.1), sy + Inches(0.05), panel_w, Inches(0.35),
             font_size=12, color=WHITE, bold=True)
    sy += Inches(0.4)
    h = Inches(0.3) * len(btns) + Inches(0.1) * (len(btns) + 1)
    add_rect(slide8, panel_x, sy, panel_w, h, WHITE)
    for j, btn in enumerate(btns):
        by = sy + Inches(0.1) + Inches(0.35) * j
        add_rect(slide8, panel_x + Inches(0.1), by, panel_w - Inches(0.2), Inches(0.3), LIGHT_BLUE)
        add_text(slide8, btn, panel_x + Inches(0.1), by + Inches(0.03), panel_w - Inches(0.2), Inches(0.25),
                 font_size=10, color=DARK, align=PP_ALIGN.CENTER)
    sy += h + Inches(0.05)


# ==================== 第9页：总结与展望 ====================
slide9 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide9, WHITE)
add_title_bar(slide9, "总结与展望")

left_x = Inches(0.85)
right_x = Inches(7.2)
y = Inches(1.6)

add_text(slide9, "已完成功能", left_x, y, Inches(5.5), Inches(0.6),
         font_size=24, color=GREEN, bold=True)

done_items = [
    "✅ 基于 YOLOv8 的目标检测模型",
    "✅ 图片检测（汽车 + 行人）",
    "✅ 视频检测（逐帧处理 + 输出）",
    "✅ 图形化操作界面（Tkinter）",
    "✅ 置信度阈值实时调节",
    "✅ 检测结果可视化与统计",
    "✅ 支持中文标签显示",
]

add_bullet_list(slide9, done_items, left_x, y + Inches(0.8), Inches(5.5), Inches(4.5),
                font_size=18, color=DARK)

add_text(slide9, "未来改进方向", right_x, y, Inches(5.3), Inches(0.6),
         font_size=24, color=ORANGE, bold=True)

future_items = [
    "🚀 模型优化：使用更大模型提升精度",
    "📷 实时摄像头：接入摄像头实时检测",
    "🚚 更多类别：扩展检测目标类别",
    "📊 数据统计：趋势分析、报表导出",
    "🌐 Web 界面：浏览器端远程访问",
    "⚡ GPU 加速：提升视频处理速度",
]

add_bullet_list(slide9, future_items, right_x, y + Inches(0.8), Inches(5.3), Inches(4.5),
                font_size=18, color=DARK)


# ==================== 第10页：致谢 ====================
slide10 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide10, WHITE)

add_rect(slide10, Inches(0), Inches(2), prs.slide_width, Inches(3.5), BLUE)
add_rect(slide10, Inches(0), Inches(5.5), prs.slide_width, Inches(0.08), DARK_BLUE)

add_text(slide10, "感谢聆听！", Inches(1), Inches(2.8), Inches(11), Inches(1.2),
         font_size=56, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

add_text(slide10, "汽车与行人检测系统", Inches(1), Inches(4.0), Inches(11), Inches(0.8),
         font_size=28, color=LIGHT_BLUE, align=PP_ALIGN.CENTER)

add_text(slide10, "Q & A", Inches(1), Inches(4.8), Inches(11), Inches(0.6),
         font_size=22, color=LIGHT_BLUE, align=PP_ALIGN.CENTER)


prs.save(output_path)
print(f"PPT 已生成: {output_path}")
print(f"共 {len(prs.slides)} 页")
